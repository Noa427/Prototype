"""
Génération de documents Word (.docx) pour compromis de vente et mandat de recherche.
Les fichiers sont retournés en bytes pour streaming HTTP.
"""
import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _header(doc: Document, title: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    doc.add_paragraph()


def _section(doc: Document, label: str, value: str):
    p = doc.add_paragraph()
    run_label = p.add_run(f"{label} : ")
    run_label.bold = True
    p.add_run(value or "—")


def generate_compromis(deal, lead) -> bytes:
    """
    Génère un compromis de vente simplifié au format .docx.
    deal: modèle Deal ORM
    lead: modèle Lead ORM
    """
    doc = Document()

    # Marges
    section = doc.sections[0]
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    _header(doc, "COMPROMIS DE VENTE")
    doc.add_paragraph(f"Date : {datetime.now().strftime('%d/%m/%Y')}")
    doc.add_paragraph()

    doc.add_heading("I. PARTIES", level=2)
    _section(doc, "Acheteur", lead.full_name)
    _section(doc, "Email acheteur", lead.email)
    _section(doc, "Téléphone", lead.phone or "—")

    doc.add_paragraph()
    doc.add_heading("II. BIEN IMMOBILIER", level=2)
    _section(doc, "Localisation", f"{deal.city or '—'} {deal.postal_code or ''}".strip())
    _section(doc, "Adresse", f"{deal.street_number or ''} {deal.street or '—'}".strip())
    _section(doc, "Type", deal.property_type or "—")
    _section(doc, "Surface", f"{deal.surface} m²" if deal.surface else "—")
    _section(doc, "DPE", deal.dpe or "—")

    doc.add_paragraph()
    doc.add_heading("III. CONDITIONS FINANCIÈRES", level=2)
    _section(doc, "Prix de vente", f"{deal.price:,} €".replace(",", " ") if deal.price else "—")
    _section(doc, "Apport personnel", f"{lead.apport:,} €".replace(",", " ") if lead.apport else "—")
    _section(doc, "Budget total acheteur", f"{lead.budget:,} €".replace(",", " ") if lead.budget else "—")
    _section(doc, "Délai envisagé", lead.delay or "—")

    doc.add_paragraph()
    doc.add_heading("IV. CONDITIONS SUSPENSIVES", level=2)
    doc.add_paragraph("• Obtention d'un prêt immobilier aux conditions habituelles du marché.")
    doc.add_paragraph("• Absence de servitudes ou hypothèques non déclarées.")
    doc.add_paragraph("• Résultats satisfaisants des diagnostics techniques obligatoires.")

    doc.add_paragraph()
    doc.add_paragraph(
        "Le présent document est un avant-contrat établi à titre indicatif. "
        "Il devra être finalisé par un notaire avant toute valeur juridique contraignante.",
        style="Normal"
    )

    doc.add_paragraph()
    doc.add_paragraph("Signature vendeur : _______________________")
    doc.add_paragraph("Signature acheteur : ______________________")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_mandat(lead, agent_name: str = "Agent AEVUM") -> bytes:
    """
    Génère un mandat de recherche simplifié au format .docx.
    lead: modèle Lead ORM
    """
    doc = Document()

    section = doc.sections[0]
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    _header(doc, "MANDAT DE RECHERCHE IMMOBILIER")
    doc.add_paragraph(f"Date : {datetime.now().strftime('%d/%m/%Y')}")
    doc.add_paragraph()

    doc.add_heading("I. MANDANT (Client)", level=2)
    _section(doc, "Nom complet", lead.full_name)
    _section(doc, "Email", lead.email)
    _section(doc, "Téléphone", lead.phone or "—")

    doc.add_paragraph()
    doc.add_heading("II. MANDATAIRE", level=2)
    _section(doc, "Agent", agent_name)
    _section(doc, "Agence", "AEVUM Immobilier")

    doc.add_paragraph()
    doc.add_heading("III. CRITÈRES DE RECHERCHE", level=2)
    _section(doc, "Budget maximum", f"{lead.budget:,} €".replace(",", " ") if lead.budget else "—")
    _section(doc, "Apport disponible", f"{lead.apport:,} €".replace(",", " ") if lead.apport else "—")
    _section(doc, "Délai d'acquisition", lead.delay or "—")
    _section(doc, "Statut du dossier", lead.status or "new")

    doc.add_paragraph()
    doc.add_heading("IV. DURÉE ET CONDITIONS", level=2)
    doc.add_paragraph(
        "Le présent mandat est consenti pour une durée de 3 mois à compter de sa signature, "
        "renouvelable par tacite reconduction."
    )
    doc.add_paragraph(
        "Le mandataire s'engage à rechercher activement tout bien correspondant aux critères "
        "définis ci-dessus et à en informer le mandant dans les meilleurs délais."
    )
    doc.add_paragraph(
        "Ce mandat est non exclusif. Le mandant conserve la liberté de rechercher par ses "
        "propres moyens ou via d'autres intermédiaires."
    )

    doc.add_paragraph()
    doc.add_paragraph("Signature mandant : _______________________")
    doc.add_paragraph("Signature mandataire : _____________________")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
